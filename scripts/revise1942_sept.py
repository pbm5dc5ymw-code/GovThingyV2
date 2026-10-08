"""Rewrite the 1942 columns from the September 1, 1942 Government Manual.

The FAQ says to use the most recent manual when a year has several, so this
replaces the March 1942 coding produced by fill1940_1942.py. It also applies
the conventions from the user's FAQ tab and feedback sheet: the Executive
Office of the President is a parent, duplicate rows share values, dead
agencies that reappear are relabeled, no intermediate parents for Social
Security Board / Fiscal Service units, related organizations belong to their
department, and quasi-official / international agencies are left blank.
"""
import openpyxl, shutil

SRC = 'data/Spreadsheet 1935-1942.xlsx'
DST = 'data/Spreadsheet 1935-1942 (Sept 1942 manual).xlsx'
COL = 33  # 1942_cabinet_agency

AG, COM, JUS, ST = 'Department of Agriculture', 'Department of Commerce', 'Department of Justice', 'Department of State'
INT, NAVY, TRE, WAR, PO = ('Department of the Interior', 'Department of the Navy', 'Department of the Treasury',
                           'Department of War', 'Post Office Department')
CABINET = {AG, COM, JUS, ST, INT, NAVY, TRE, WAR, PO, 'Department of Labor'}
EOP, OEM, FSA, FWA = 'Executive Office of the President', 'Office for Emergency Management', 'Federal Security Agency', 'Federal Works Agency'
NHA, FHLBA, FPHA = 'National Housing Agency', 'Federal Home Loan Bank Administration', 'Federal Public Housing Authority'
ARA, ACAA, SOS, AGF = ('Agricultural Research Administration', 'Agricultural Conservation and Adjustment Administration',
                       'Services of Supply', 'Army Ground Forces')
RFC, CSC = 'Reconstruction Finance Corporation', 'United States Civil Service Commission'
P1, P2 = 'Office of the First Assistant Postmaster General', 'Office of the Second Assistant Postmaster General'
DEAD, OTHER, SKIP = 'DEAD', 'OTHER', 'SKIP'


def vals(c):
    if c == DEAD: return ['Dead'] * 5
    if c == OTHER: return ['No', 'No', 'None', 'None', 'Yes']
    if c in CABINET: return ['No', 'Yes', c, 'None', 'No']
    return ['No', 'Yes', 'None', c, 'No']


def rows(code, *rs):
    return {r: code for r in rs}


O = {}
# Executive Office of the President as parent
O.update(rows(OTHER, 479))
O.update(rows(EOP, 480, 13, 374, 377, 381))
O.update(rows(DEAD, 376, 467))                          # Office of Government Reports, USIS -> OWI (June 1942)
# Office for Emergency Management units
O.update(rows(OEM, 509, 488, 494, 496, 498, 499, 508, 506))  # OCD (both rows), OWI, WMC, WSA, WRA, Alien Property Custodian, BWC
O.update(rows(DEAD, 490, 529, 528, 500))                # DCB -> BWC, Facts and Figures, Coordinator of Information, Defense Housing Coordination
# Emergency war agencies outside OEM
O.update(rows(OTHER, 511, 512, 514, 516, 515, 513))     # OPA, BEW, Censorship, Selective Service, War Relief Control Board, NHA
O.update(rows(NHA, 321, 324, 325))                      # FHLB Administration, FHA, FPHA
O.update(rows(FHLBA, 322, 336, 349, 323))               # FHLB Board, FSLIC, HOLC, FHLB System
O.update(rows(FPHA, 437))                               # Defense Homes Corporation
O.update(rows(DEAD, 351, 328))                          # USHA -> FPHA; Federal Loan Agency dissolved
# RFC group -> Commerce (EO 9071, Feb 24, 1942)
O.update(rows(COM, 432, 304, 297, 300, 329, 440))
O.update(rows(RFC, 44))                                 # War Damage Corporation
O.update(rows(DEAD, 69))                                # Bureau of Marine Inspection and Navigation split up
# Agriculture (EO 9069, Feb 23, 1942)
O.update(rows(AG, 30, 35, 59))                          # Agricultural Marketing Admin, Agricultural Research Admin, Agricultural War Relations
O.update(rows(ARA, 34, 36, 39, 40, 41, 42, 51))
O.update(rows(ACAA, 4, 54, 45, 55))                     # Agricultural Adjustment Agency, SCS, FCIC, Sugar Agency
O.update(rows(DEAD, 29, 57, 339, 504))                  # merged into AMA; Agricultural Defense Relations renamed
O.update(rows('Federal Deposit Insurance Corporation', 314))  # Federal Credit Union System
# Interior
O.update(rows(INT, 194, 208, 200, 370, 369))            # Fishery Coordination, War Resources Council, Solid Fuels for War, NPPC, Park Trust Fund
O.update(rows(DEAD, 532, 201, 187, 309))                # renamed "for War"; Division of Investigations; Board of Surveys and Maps
# Navy
O.update(rows(NAVY, 232))                               # Bureau of Naval Personnel
O.update(rows(DEAD, 222))
# Treasury
O.update(rows(TRE, 244, 235, 243, 261, 260))            # War Savings Staff; Fiscal Service units directly under Treasury
O.update(rows(DEAD, 530))
# War (reorganization of March 9, 1942)
O.update(rows(WAR, 271, 286))
O.update(rows(SOS, 274, 294, 539, 277, 278, 276, 536, 291, 293, 296, 279, 280, 292, 281))
O.update(rows(AGF, 537))
O.update(rows(DEAD, 287, 288, 289, 290, 295, 538))
# Justice, Labor, State
O.update(rows(JUS, 89, 91, 102))
O.update(rows(DEAD, 75, 311))
O.update(rows(ST, 111, 129, 157))
O.update(rows(DEAD, 534, 535))
# Post Office: Offices renamed Bureaus; both rows of each pair share values
O.update(rows(PO, 388, 394, 401, 411, 419, 422))
O[399] = P1                                             # Division of Rural Mails
O[398] = P2                                             # Division of Air Mail Service
# Commerce related organizations, FSA, duplicates
O.update(rows(COM, 452, 341))
O.update(rows(FSA, 484, 485, 486, 487, 46, 245, 518, 20, 517))
# Corrections after review against page layout and the feedback sheet
O.update(rows('Bureau of Accounts', 408, 423))          # indented under Bureau of Accounts (p. 224)
O.update(rows(RFC, 297, 329, 440))                      # RFC subsidiaries moved with RFC to Commerce
O.update(rows(OEM, 511, 512, 513))                      # OPA, BEW, National Housing Agency
O.update(rows(EOP, 514, 515, 516))                      # Censorship, War Relief Control Board, Selective Service
O.update(rows(NHA, 322, 323, 336, 349))                 # National Housing Agency units, no intermediate parent
O.update(rows(CSC, 26, 519))
O.update(rows(FWA, 344, 477))
O.update(rows(SKIP, 6, 358, 383, 382))                  # quasi-official / international

NEW = [
    ('Agricultural Conservation and Adjustment Administration', AG),
    ('Office of Petroleum Coordinator for War', INT),
    ('Office of Strategic Services', OTHER),
    ('American Hemisphere Exports Office', ST),
]

shutil.copy(SRC, DST)
wb = openpyxl.load_workbook(DST)
ws = wb.worksheets[0]
assert ws.cell(2, COL).value == '1942_cabinet_agency'

for r, c in O.items():
    for i in range(5):
        ws.cell(r, COL + i).value = None if c == SKIP else vals(c)[i]
    if c not in (DEAD, SKIP) and not (ws.cell(r, 2).value or '').strip():
        ws.cell(r, 2).value = 'Executive'

r = max(x for x in range(3, ws.max_row + 1) if ws.cell(x, 1).value) + 1
for name, c in NEW:
    ws.cell(r, 1).value = name
    ws.cell(r, 2).value = 'Executive'
    for col in range(3, COL):
        ws.cell(r, col).value = ['No', 'No', 'None', 'None', 'No'][(col - 3) % 5]
    for i, v in enumerate(vals(c)):
        ws.cell(r, COL + i).value = v
    r += 1

wb.save(DST)
print('saved', DST)
