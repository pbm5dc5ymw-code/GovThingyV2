import openpyxl, shutil

SRC = 'data/Spreadsheet of 1935 and 1936 and 1937.xlsx'
DST = 'data/Spreadsheet of 1935 1936 1937 and 1939.xlsx'

AG, COM, JUS, LAB, ST = 'Department of Agriculture', 'Department of Commerce', 'Department of Justice', 'Department of Labor', 'Department of State'
INT, NAVY, TRE, WAR, PO = 'Department of the Interior', 'Department of the Navy', 'Department of the Treasury', 'Department of War', 'Post Office Department'
CABINET = {AG, COM, JUS, LAB, ST, INT, NAVY, TRE, WAR, PO}

DEAD = 'DEAD'
NOTYET = 'NOTYET'
CAB = 'CAB'
OTHER = 'OTHER'
SKIP = 'SKIP'  # international / not considered: leave 1939 blank

def vals(c):
    if c == DEAD: return ['Dead'] * 5
    if c == NOTYET: return ['No', 'No', 'None', 'None', 'No']
    if c == CAB: return ['Yes', 'No', 'None', 'None', 'No']
    if c == OTHER: return ['No', 'No', 'None', 'None', 'Yes']
    # child of a parent
    if c in CABINET: return ['No', 'Yes', c, 'None', 'No']
    return ['No', 'Yes', 'None', c, 'No']

FSA, FWA, FLA, EOP = 'Federal Security Agency', 'Federal Works Agency', 'Federal Loan Agency', 'Executive Office of the President'
FCA, FHLBB, FRS, SSB = 'Farm Credit Administration', 'Federal Home Loan Bank Board', 'Federal Reserve System', 'Social Security Board'
OGR, NARA = 'Office of Government Reports', 'National Archives'
P1, P2, P3, P4 = ('Office of the First Assistant Postmaster General', 'Office of the Second Assistant Postmaster General',
                  'Office of the Third Assistant Postmaster General', 'Office of the Fourth Assistant Postmaster General')

# Row number -> 1939 classification (from GPO-GOVMAN-1939-10-01)
M = {}
def s(rows, c):
    for r in rows: M[r] = c

# Executive Office of the President (created 1939)
s([479], OTHER); s([480, 13, 374, 376, 377], EOP)
s([14], 'Bureau of the Budget'); s([467], OGR)
s([26, 519, 22, 23], OTHER)
# Cabinet departments
s([28, 58, 72, 97, 109, 171, 213, 233, 268, 387], CAB)
# Agriculture
s([3, 4, 29, 31, 32, 34, 36, 39, 40, 41, 42, 45, 46, 47, 48, 50, 51, 52, 53, 54, 55, 56,
   24, 305, 306, 338, 441, 478], AG)
s([33, 38, 43], DEAD)
s([9, 302, 303, 314, 319, 326, 327, 428, 438], FCA)
# Commerce
s([60, 61, 63, 64, 66, 67, 69, 352], COM)
s([65], DEAD)
# Justice
s([74, 75, 76, 77, 78, 79, 80, 81, 82, 84, 86, 87, 88, 90, 91, 92], JUS)
s([93], DEAD)
# Labor
s([85, 98, 99, 100, 101, 102, 104, 107, 108], LAB)
s([103, 105, 311], DEAD)
# State
s([110, 112, 113, 114, 117, 118, 119, 123, 124, 125, 126, 132, 133, 137, 138, 139, 140, 143, 144,
   146, 147, 148, 152, 153, 156, 158, 159, 160, 161, 166, 167, 168, 169, 170], ST)
s([127, 162, 164], DEAD)
# Interior
s([62, 37, 172, 173, 174, 176, 177, 178, 180, 181, 182, 185, 186, 187, 189, 191, 192, 195, 197, 199,
   203, 207, 209, 210, 211, 212, 309, 357, 363], INT)
s([188, 362], DEAD)
# Navy
s([215, 218, 219, 220, 221, 222, 223, 224, 225, 227, 228, 229, 230, 231], NAVY)
# Treasury
s([234, 236, 237, 238, 239, 240, 241, 246, 247, 248, 250, 252, 254, 255, 258, 259, 260, 261, 262,
   264, 265, 266, 267], TRE)
s([474], DEAD)
# War
s([269, 274, 275, 276, 277, 278, 279, 280, 287, 288, 289, 290, 291, 292, 293, 294, 295, 296,
   455, 456, 457, 462], WAR)
s([282], DEAD)
# Post Office
s([389, 393, 400, 410, 419, 420, 421, 422, 424, 427], PO)
s([390, 391, 392, 398], P1); s([395, 396, 397, 399], P2)
s([402, 403, 404, 405, 406, 407, 408, 409], P3); s([412, 413, 414, 415, 416, 417], P4)
s([425], 'Bureau of Accounts')
# Federal Security Agency (created 1939)
s([337], OTHER); s([20, 517, 375, 193, 245, 518, 447], FSA); s([484, 485, 486, 487], SSB)
# Federal Works Agency (created 1939)
s([343], OTHER); s([317, 320, 332, 344, 351, 429, 430, 477], FWA)
s([316, 476], DEAD)
# Federal Loan Agency (created 1939)
s([328], OTHER); s([297, 300, 304, 322, 324, 329, 432, 440], FLA)
s([323, 335, 336, 349], FHLBB); s([442], DEAD)
# Federal Reserve / Archives
s([10, 334], OTHER); s([307, 333], FRS); s([361, 366], NARA)
# Independent agencies still listed
s([5, 6, 8, 17, 21, 308, 312, 315, 330, 331, 340, 341, 342, 354, 356, 358, 359, 360, 364, 367, 368,
   369, 370, 431, 445, 446, 451, 452, 453, 460, 461, 466, 469, 471, 473, 475], OTHER)
# Abolished / not listed in 1939
s([318, 365, 454, 470], DEAD)
# International (not considered)
s([382, 383, 481, 482, 483], SKIP)

NEW = [
    ('Executive Committee on Commercial Policy', OTHER),
    ('Office of the Solicitor General', JUS),
    ('Joint Army and Navy Munitions Board', WAR),
    ('Extension Service', AG),
    ('Office of Information (Department of Agriculture)', AG),
    ('Business Advisory Council', COM),
    ('Bureau of Old-Age and Survivors Insurance', SSB),
    ('Committee on Economic Security', OTHER),
    ('Interdepartmental Committee to Coordinate Health and Welfare Activities', OTHER),
    ('Division of Press Intelligence', OGR),
    ('Air Safety Board', 'Civil Aeronautics Authority'),
    ('Federal Open Market Committee', FRS),
    ('National Research Council', 'National Academy of Sciences'),
    ("Board of Veterans' Appeals", 'Veterans Administration'),
]

shutil.copy(SRC, DST)
wb = openpyxl.load_workbook(DST)
ws = wb.worksheets[0]
C39 = 18  # column R = 1939_cabinet_agency
assert ws.cell(2, C39).value == '1939_cabinet_agency'

auto_dead, auto_notyet = [], []
for r in range(3, 520):
    name = (ws.cell(r, 1).value or '').strip()
    y37 = [str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else '' for c in range(13, 18)]
    if r in M:
        c = M[r]
    elif not name or all(x == '' for x in y37):
        continue  # blank / legislative / judicial rows stay blank
    elif 'Dead' in y37:
        c = DEAD
    elif y37 == ['No', 'No', 'None', 'None', 'No']:
        c = NOTYET; auto_notyet.append((r, name))
    else:
        raise SystemExit(f'Unclassified live row {r}: {name} {y37}')
    if c == SKIP:
        continue
    for i, v in enumerate(vals(c)):
        ws.cell(r, C39 + i).value = v
    if c not in (DEAD, NOTYET) and not (ws.cell(r, 2).value or '').strip():
        ws.cell(r, 2).value = 'Executive'

r = 520
for name, c in NEW:
    ws.cell(r, 1).value = name
    ws.cell(r, 2).value = 'Executive'
    for col in range(3, 18):  # 1935-1937: not yet listed
        ws.cell(r, col).value = ['No', 'No', 'None', 'None', 'No'][(col - 3) % 5]
    for i, v in enumerate(vals(c)):
        ws.cell(r, C39 + i).value = v
    r += 1

wb.save(DST)
print('saved', DST)
print('auto not-yet rows kept as not listed:', len(auto_notyet))
