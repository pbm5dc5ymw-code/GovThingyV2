"""Fill the 1946 columns from the 1946 Government Manual, First Edition (the only 1946 edition on govinfo).

Organization charts are back in this edition (printed pp. 569-602), so parents come first from the charts
(solid line = authority, dashed = not a parent), then from unit lists, officials pages and the Executive
Office (p. 52) and OEM (p. 59) membership lists. Every row was run through
skills/gov-manual-agency-entry/scripts/evidence.py; a row stays live only with printed evidence.
Page numbers are printed pages (PDF page = printed + 10).
"""
import openpyxl, shutil

SRC = 'data/Spreadsheet 1935-1945.xlsx'
DST = 'data/Spreadsheet 1935-1946.xlsx'
YEAR, PREV_YEAR = 1946, 1945

AG, COM, JUS, LAB, ST = ('Department of Agriculture', 'Department of Commerce', 'Department of Justice',
                         'Department of Labor', 'Department of State')
INT, NAVY, TRE, WAR, PO = ('Department of the Interior', 'Department of the Navy', 'Department of the Treasury',
                           'Department of War', 'Post Office Department')
CABINET = {AG, COM, JUS, LAB, ST, INT, NAVY, TRE, WAR, PO}
DEAD, NOTYET, OTHER, CAB = 'DEAD', 'NOTYET', 'OTHER', 'CAB'
OEM, RFC, FSA, FWA = 'Office for Emergency Management', 'Reconstruction Finance Corporation', 'Federal Security Agency', 'Federal Works Agency'
BFDC, NIA, SS = 'Bureau of Foreign and Domestic Commerce', 'National Intelligence Authority', 'Special Staff'
B1 = 'Bureau of the First Assistant Postmaster General'
# State offices (chart p. 576; officials pp. 129-132)
ESP, FL, BF = 'Office of Economic Security Policy', 'Office of Foreign Liquidation', 'Office of Budget and Finance (State)'
ORI, OICD = 'Office of Research and Intelligence', 'Office of Intelligence Collection and Dissemination'
IICA, DA, CTL = 'Office of International Information and Cultural Affairs', 'Office of Departmental Administration', 'Office of Controls'


def vals(c):
    """Five-column coding: cabinet, child, cabinet parent, other parent, other agency."""
    if c == DEAD: return ['Dead'] * 5
    if c == NOTYET: return ['No', 'No', 'None', 'None', 'No']
    if c == CAB: return ['Yes', 'No', 'None', 'None', 'No']
    if c == OTHER: return ['No', 'No', 'None', 'None', 'Yes']
    if c in CABINET: return ['No', 'Yes', c, 'None', 'No']
    return ['No', 'Yes', 'None', c, 'No']


O = {}
def rows(code, *row_numbers):
    O.update({r: code for r in row_numbers})

O = {}
def rows(code, *row_numbers):
    O.update({r: code for r in row_numbers})

# Executive Office (p. 52, chart p. 571) and emergency agencies (OEM list p. 59)
rows(DEAD, 562)                                 # War Refugee Board terminated Sept 14, 1945 (App. A)
rows(DEAD, 380)          # Foreign Economic Administration terminated Sept 27, 1945 (App. A)
rows(DEAD, 494, 496, 497, 507, 514, 542)        # OWI, WMC, WPB, National War Labor Board, Censorship, OSS terminated 1945 (App. A)
rows(DEAD, 545, 548, 567, 569)                  # Training Within Industry (WMC ended); Smaller War Plants Corp. (functions to RFC/Commerce); Petroleum Reserves Corp. (became War Assets Corp.); Committee on Medical Research (one sentence, p. 75)
rows(DEAD, 647)                                 # Surplus Property Board: functions merged into War Assets Corporation Feb 1, 1946 (p. 105)
rows(RFC, 566, 568)                             # U. S. Commercial Company, Rubber Development Corporation under RFC (FLA chart p. 584)
rows(LAB, 564, 105, 311)                        # Retraining and Reemployment Administration, U. S. Employment Service, Apprentice Training Service (Labor chart p. 583)
# Commerce (p. 317)
rows(DEAD, 649, 650)                            # Office of Civilian Defense Property and Office of Surplus Property not in list; OSP terminated Oct 1945 (App. A)
# Justice (chart p. 579)
rows(DEAD, 89)                                  # War Division not on chart or list
# Interior (p. 270 list)
rows(DEAD, 194, 197)                            # Office of Fishery Coordination terminated Oct 29, 1945 (App. A); Office of the Chief Clerk not in list
# Treasury (chart p. 577, list p. 159)
rows(TRE, 267)                                  # Coast Guard returned to Treasury Jan 1, 1946
# War (chart p. 578; Special Staff of 12 divisions p. 197)
rows(WAR, 270)                                  # Special Staff (chart box)
rows(SS, 275, 283, 294, 539)                    # Inspector General, Bureau of Public Relations, National Guard Bureau, Reserve and ROTC Affairs listed in the Special Staff
rows(DEAD, 456, 281)                            # Joint Economy Board and Morale Branch not mentioned
# Post Office (p. 221)
rows(B1, 391)                                   # Division of Postmasters reappears under the First Assistant (FAQ: unhide)
# State (chart p. 576; officials pp. 129-132)
rows(DEAD, 657, 665, 624, 671, 672, 622)        # Interim Int'l Information Service, War Areas Economic, Geography and Cartography, Cultural Cooperation, International Information, Budget and Finance divisions not listed
rows(ESP, 669)                                  # Division of Economic Security Controls under Office of Economic Security Policy
rows(DA, 620, 123)                              # Central Translating Division; Division of Communications and Records reappears (FAQ: unhide)

NEW = [
    ('Civilian Production Administration', OEM),                    # p. 59; successor to WPB Nov 3, 1945
    ('Director of Liquidation', OEM),                               # p. 63
    ('War Assets Administration', OEM),                             # p. 75
    (NIA, OTHER),                                                   # p. 85
    ('Central Intelligence Group', NIA),                            # p. 85 subheading
    ('National Wage Stabilization Board', LAB),                     # Labor chart p. 583; successor to NWLB
    ('Wage Adjustment Board', LAB),                                 # chart p. 583
    ('Shipbuilding Stabilization Committee', LAB),                  # chart p. 583
    ('Office of the Chief Clerk and Budget Officer (Labor)', LAB),  # chart p. 583
    ('Office of the Director of Personnel (Labor)', LAB),
    ('Office of the Director of Information (Labor)', LAB),
    ('Office of the Solicitor (Labor)', LAB),
    ('Library (Labor)', LAB),
    ('Division of Public Relations (Treasury)', TRE),               # chart p. 577
    ('United States Attorneys', 'Administrative Division'),         # chart p. 579
    ('United States Marshals', 'Administrative Division'),          # chart p. 579
    ('Strength Accounting and Statistics Office', WAR),             # chart p. 578
    ('Personnel and Administrative Branch (War)', WAR),             # chart p. 578
    ('Secretariat (War)', WAR),                                     # chart p. 578
    ('Office of the Secretary (Interior)', INT),                    # list p. 270
    ('Division of Administrative Services (Interior)', INT),
    ('Office of Labor Relations (Interior)', INT),
    ('Office of Field Representatives (Interior)', INT),
    ('Office of Small Business', BFDC),                             # list p. 317 (indented under BFDC)
    ('Office of Business Economics', BFDC),
    ('Office of Field Operations', BFDC),
    ('Office of Domestic Commerce', BFDC),
    ('Office of International Trade', BFDC),
    ("St. Elizabeths Hospital", FSA),                               # FSA chart p. 585 (solid line)
    ('Office of War Property Distribution', FSA),                   # chart p. 585
    ('Office of General Counsel (FWA)', FWA),                       # FWA chart p. 586
    ('Office of Executive Officer (FWA)', FWA),
    ('Office of Economic Research (FWA)', FWA),
    ('Office of Information (FWA)', FWA),
    # State
    (ESP, ST), (FL, ST), (BF, ST), (ORI, ST), (OICD, ST),
    ('Division of German and Austrian Economic Affairs', ESP),
    ('Division of Japanese and Korean Economic Affairs', ESP),
    ('Division of Budget', BF), ('Division of Finance (State)', BF), ('UNRRA Division', BF),
    ('Division of Investigations (State)', CTL),
    ('International Press and Publications Division', IICA),
    ('International Broadcasting Division', IICA),
    ('International Motion Pictures Division', IICA),
    ('Division of International Exchange of Persons', IICA),
    ('Division of Libraries and Institutes', IICA),
    ('Area Division I (Europe)', IICA), ('Area Division II (Near East and Africa)', IICA),
    ('Area Division III (Far East)', IICA), ('Area Division IV (American Republics)', IICA),
    ('Area Division V (Occupied Area)', IICA),
    ('Division of American Republics Intelligence', ORI),
    ('Division of British Commonwealth Intelligence', ORI),
    ('Division of Europe, Near East and Africa Intelligence', ORI),
    ('Division of Far East Intelligence', ORI),
    ('Division of USSR Intelligence', ORI),
    ('Division of International and Functional Intelligence', ORI),
    ('Division of Map Intelligence and Cartography', ORI),
    ('Division of Intelligence Acquisition and Distribution', OICD),
    ('Intelligence Reference Division', OICD),
    ('Division of Biographic Intelligence', OICD),
    ('Presentation Division', OICD),
]

shutil.copy(SRC, DST)
wb = openpyxl.load_workbook(DST)
ws = wb.worksheets[0]
hdr = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value}
COL, PREV = hdr[f'{YEAR}_cabinet_agency'], hdr[f'{PREV_YEAR}_cabinet_agency']

last = max(r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value)
for r in range(3, last + 1):
    if not ws.cell(r, 1).value:
        continue
    prev = [ws.cell(r, PREV + i).value for i in range(5)]
    if r in O:
        new = vals(O[r])
        if O[r] not in (DEAD, NOTYET) and not (ws.cell(r, 2).value or '').strip():
            ws.cell(r, 2).value = 'Executive'
    elif all(v in (None, '') for v in prev):
        continue
    else:
        new = prev
    for i, v in enumerate(new):
        ws.cell(r, COL + i).value = v

r = last + 1
for name, c in NEW:
    ws.cell(r, 1).value = name
    ws.cell(r, 2).value = 'Executive'
    for col in range(3, COL):
        ws.cell(r, col).value = vals(NOTYET)[(col - 3) % 5]
    for i, v in enumerate(vals(c)):
        ws.cell(r, COL + i).value = v
    r += 1

# FAQ: hide (never delete) dead rows and legislative/judicial rows
for x in range(3, ws.max_row + 1):
    if ws.cell(x, 1).value:
        kind = (ws.cell(x, 2).value or '').strip()
        ws.row_dimensions[x].hidden = kind in ('Legislative', 'Judicial') or ws.cell(x, COL).value == 'Dead'

wb.save(DST)
print('saved', DST, 'new rows', last + 1, '-', r - 1)
