"""Fill the 1948 columns from the 1948 Government Manual (the only 1948 edition on govinfo).

Organization charts (printed pp. 573-610) decide parents first, then unit lists, officials pages, the
Executive Office chart (p. 575) and OEM list (p. 65), and Appendix A (Reorganization Plans 1 and 3 of
1947, National Security Act of 1947). Every row was run through evidence.py. Rows are found by name.
Page numbers are printed pages (PDF page = printed + 8).

National Security Act of 1947: the National Military Establishment (Secretary of Defense, a Cabinet
member) is coded as the cabinet department; the Departments of the Army, Navy and Air Force are coded as
its children (their Secretaries left the Cabinet). Units of the Army, Navy and Air Force therefore name
their department in the other_parent column. OPEN convention - flagged for the user.
"""
import openpyxl, shutil

SRC = 'data/Spreadsheet 1935-1947.xlsx'
DST = 'data/Spreadsheet 1935-1948.xlsx'
YEAR, PREV_YEAR = 1948, 1947

AG, COM, JUS, LAB, ST = ('Department of Agriculture', 'Department of Commerce', 'Department of Justice',
                         'Department of Labor', 'Department of State')
INT, TRE, PO = 'Department of the Interior', 'Department of the Treasury', 'Post Office Department'
NME, ARMY, NAVY, AF = 'National Military Establishment', 'Department of the Army', 'Department of the Navy', 'Department of the Air Force'
WAR = 'Department of War'
CABINET = {AG, COM, JUS, LAB, ST, INT, TRE, PO, NME}
DEAD, NOTYET, OTHER, CAB = 'DEAD', 'NOTYET', 'OTHER', 'CAB'
EOP, OEM, FSA, HHFA, HLBB = 'Executive Office of the President', 'Office for Emergency Management', 'Federal Security Agency', 'Housing and Home Finance Agency', 'Home Loan Bank Board'
SS, JCS, NSC, PHS = 'Special Staff', 'Joint Chiefs of Staff', 'National Security Council', 'Public Health Service'
UNA, OII, OEE, BP = 'Office of United Nations Affairs', 'Office of International Information', 'Office of Educational Exchange', 'Office of Budget and Finance (State)'
FDP, NEA, CTL = 'Office of Financial and Development Policy', 'Office of Near Eastern and African Affairs', 'Office of Controls'


def vals(c):
    """Five-column coding: cabinet, child, cabinet parent, other parent, other agency."""
    if c == DEAD: return ['Dead'] * 5
    if c == NOTYET: return ['No', 'No', 'None', 'None', 'No']
    if c == CAB: return ['Yes', 'No', 'None', 'None', 'No']
    if c == OTHER: return ['No', 'No', 'None', 'None', 'Yes']
    if c in CABINET: return ['No', 'Yes', c, 'None', 'No']
    return ['No', 'Yes', 'None', c, 'No']


O = {}
NAMED = []
def rows(code, *row_numbers):
    O.update({r: code for r in row_numbers})
def named(code, *names):
    NAMED.extend((n, code) for n in names)

# Executive Office (chart p. 575) and emergency agencies (OEM list p. 65)
named(DEAD, 'Office of Government Reports', 'Office of Scientific Research and Development', 'National Defense Research Committee',
      'Price Decontrol Board', 'Office of Selective Service Records', 'Federal Board of Hospitalization')  # not on EOP chart; OSRD ended Dec 31, 1947; Selective Service Records absorbed; FBH terminated June 30, 1948 (App. A)
named(OTHER, 'War Assets Administration', 'Selective Service System')   # WAA not in OEM list (p. 65); Selective Service System re-created by act of June 24, 1948 (p. 466; FAQ: unhide)
# National Military Establishment (chart p. 582)
named(DEAD, WAR, 'The Army Air Forces', 'The Joint Board', 'The Joint Economy Board', 'War Department Manpower Board', 'Budget Division (War)')  # War renamed Department of the Army; AAF became Department of the Air Force; joint boards abolished 1947 (App. A)
named(NME, 'Department of the Navy', JCS)
named(JCS, 'Army and Navy Staff College')        # now The National War College (p. 162)
# Army (chart p. 583): every unit carried from War is re-parented below; these are on the chart
named(ARMY, 'War Department General Staff', 'Special Staff', 'Office of The Adjutant General', 'Office of the Chief of Chaplains',
      'Office of the Chief of Finance', 'Office of the Chief of Engineers', 'Office of the Chief Signal Officer', 'Office of the Quartermaster General',
      'Office of the Surgeon General', 'Office of the Chief of Ordnance', 'Office of the Chief of the Chemical Warfare Service',
      'Office of the Provost Marshal General', 'Office of the Chief of Transportation', 'Special Services (War)', 'Army Ground Forces',
      'The Panama Canal', 'Arlington Memorial Ampitheater Commission')
rows(ARMY, 276)                                  # Judge Advocate General's Department (Army row)
# State (chart p. 580; officials pp. 83-90)
named(DEAD, 'Office of Special Political Affairs', 'Division of International Organization Affairs',
      'Office of International Information and Cultural Affairs', 'Area Division I (Europe)', 'Area Division II (Near East and Africa)',
      'Area Division III (Far East)', 'Area Division IV (American Republics)', 'Area Division V (Occupied Area)',
      'Office of Economic Security Policy', 'Division of Economic Security Controls', 'Division of German and Austrian Economic Affairs',
      'Division of Japanese and Korean Economic Affairs', 'Division of Lend-Lease and Surplus War Property Affairs',
      'Division of Foreign Economic Development', 'UNRRA Division', 'Division of Budget', 'Division of Management Planning',
      'Division of Middle Eastern Affairs', 'Special War Problems Division')   # Special Projects Division not in Office of Controls list
named(UNA, 'Division of International Security Affairs', 'Division of Dependent Area Affairs')
named(OII, 'International Press and Publications Division', 'International Broadcasting Division', 'International Motion Pictures Division')
named(OEE, 'Division of International Exchange of Persons', 'Division of Libraries and Institutes')
# Treasury (list p. 102, chart p. 581)
named(DEAD, 'Division of Monetary Research', 'Division of Research and Statistics', 'Office of Chief Clerk',
      'Office of Superinendent of Treasury Buildings', 'Foreign Funds Control')
# Interior (list p. 230, chart p. 588)
named(DEAD, 'Coal Mines Administration', 'Solid Fuels Administration for War', 'Office of the Safety Engineer (Interior)',
      'Office of Labor Relations (Interior)', 'Office of Field Representatives (Interior)')   # SFAW terminated June 30, 1947 (App. A)
# Agriculture (chart p. 589)
named(DEAD, 'Sugar Rationing Administration')
# Commerce (chart p. 590)
named(DEAD, 'Office of Small Business')
# Labor (list p. 318, chart p. 591)
named(DEAD, 'United States Conciliation Service')  # functions to Federal Mediation and Conciliation Service (independent, p. 345)
# Federal Security Agency (chart p. 597): Freedmen's Hospital drawn inside the Public Health Service box
named(PHS, "Freedmen's Hospital")
# Housing: National Housing Agency dissolved July 27, 1947; Housing and Home Finance Agency (chart p. 601)
named(DEAD, 'National Housing Agency', 'Federal Home Loan Bank Administration', 'Federal Public Housing Authority')
named(HHFA, 'Federal Housing Administration')
named(HLBB, 'Federal Home Loan Bank System', "Home Owners' Loan Corporation", 'Federal Savings and Loan Insurance Corporation')
# Federal Loan Agency abolished June 30, 1947 (App. A); RFC independent; subsidiaries dissolved or only run-in mentions
named(DEAD, 'Federal Loan Agency', 'RFC Mortgage Company', 'War Damage Corporation', 'U. S. Commercial Company', 'Rubber Development Corporation')
named(OTHER, 'Reconstruction Finance Corporation')

NEW = [
    (NME, CAB),                                                     # p. 141, chart p. 582
    (ARMY, NME), (AF, NME),                                         # charts pp. 583, 585
    ('War Council', NME), ('Munitions Board', NME), ('Research and Development Board', NME),
    ('Joint Staff', JCS),
    ('Industrial College of the Armed Forces', JCS),                # p. 162
    ('Aeronautical Board', NME), ('Armed Services Petroleum Board', NME), ('Army-Navy Explosives Safety Board', NME),  # pp. 158-160
    (NSC, EOP), ('Central Intelligence Agency', NSC),               # p. 61
    ('National Security Resources Board', EOP),                     # p. 62
    ('Economic Cooperation Administration', OTHER),                 # p. 68
    ('Federal Mediation and Conciliation Service', OTHER),          # p. 345
    (HHFA, OTHER), (HLBB, HHFA), ('Public Housing Administration', HHFA),   # chart p. 601
    ('Loyalty Review Board', 'United States Civil Service Commission'),     # p. 499
    # State
    ('Policy Planning Staff', ST), ('Coordinator for Foreign Aid and Assistance', ST), (UNA, ST), (OII, ST), (OEE, ST),
    ('Division of United Nations Economic Affairs', UNA), ('Division of United Nations Political Affairs', UNA),
    ('Division of Greek, Turkish, and Iranian Affairs', NEA), ('Division of South Asian Affairs', NEA),
    ('Division of Economic-Property Policy', FDP), ('Division of Occupied-Area Economic Affairs', FDP),
    ('Division of Organization and Budget', BP), ('Division of Procurement Control', BP),
    ('Division of Protective Services', CTL),
    # Treasury (list p. 102)
    ('Office of International Finance', TRE), ('Office of the Technical Staff', TRE), ('Office of Administrative Services (Treasury)', TRE),
    # Interior (chart p. 588)
    ('Division of Programs (Interior)', INT),
    # Agriculture (chart p. 589)
    ('Administrator, Research and Marketing Act', AG),
    # Labor (list p. 318)
    ('Office of International Labor Affairs', LAB), ("Bureau of Veterans' Reemployment Rights", LAB),
    # Federal Security Agency (chart p. 597)
    ('Office of Publications and Reports (FSA)', FSA),
    ('National Institute of Health', PHS), ('Bureau of Medical Services', PHS), ('Bureau of State Services', PHS),
]

shutil.copy(SRC, DST)
wb = openpyxl.load_workbook(DST)
ws = wb.worksheets[0]
hdr = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value}
COL, PREV = hdr[f'{YEAR}_cabinet_agency'], hdr[f'{PREV_YEAR}_cabinet_agency']

last = max(r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value)
byname = {}
for r in range(3, last + 1):
    if ws.cell(r, 1).value:
        byname.setdefault(str(ws.cell(r, 1).value).strip(), []).append(r)
for n, c in NAMED:
    assert n in byname, n
    for r in byname[n]:
        O[r] = c
for r in range(3, last + 1):
    if not ws.cell(r, 1).value:
        continue
    prev = [ws.cell(r, PREV + i).value for i in range(5)]
    if r in O:
        new = vals(O[r])
        if O[r] not in (DEAD, NOTYET) and not (ws.cell(r, 2).value or '').strip():
            ws.cell(r, 2).value = 'Executive'
    elif all(v in (None, '') for v in prev):
        continue
    else:
        new = prev
        cp = str(prev[2] or '').strip()
        if cp in (WAR, NAVY):           # Army/Navy are no longer cabinet departments: parent moves to other_parent
            new = vals(ARMY if cp == WAR else NAVY)
    for i, v in enumerate(new):
        ws.cell(r, COL + i).value = v

r = last + 1
for name, c in NEW:
    ws.cell(r, 1).value = name
    ws.cell(r, 2).value = 'Executive'
    for col in range(3, COL):
        ws.cell(r, col).value = vals(NOTYET)[(col - 3) % 5]
    for i, v in enumerate(vals(c)):
        ws.cell(r, COL + i).value = v
    r += 1

# FAQ: hide (never delete) dead rows and legislative/judicial rows
for x in range(3, ws.max_row + 1):
    if ws.cell(x, 1).value:
        kind = (ws.cell(x, 2).value or '').strip()
        ws.row_dimensions[x].hidden = kind in ('Legislative', 'Judicial') or ws.cell(x, COL).value == 'Dead'

wb.save(DST)
print('saved', DST, 'new rows', last + 1, '-', r - 1)
