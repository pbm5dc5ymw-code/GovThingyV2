"""Fill the 1950 columns from the United States Government Organization Manual 1950-51 (only edition).

Organization charts (printed pp. 492-530) decide parents first: solid line = parent, dashed/dotted =
not a parent; units printed inside a box headed by an official (e.g. "Assistant Chief of Staff, G-4") go to
the department, boxes naming an organizational unit (Bureau, Office, Service) are parents. Then unit
lists, officials pages, section headings ("RELATED ORGANIZATIONS" subheadings) and Appendix A
(Reorganization Plans of 1949 and 1950). Every row was run through evidence.py.
Page numbers are printed pages (PDF page = printed + 8).

Renames keep the same row unless the sheet already has a row with the new name (codebook rule):
  National Military Establishment -> Department of Defense (row 803, executive department Aug 10, 1949);
  War Council -> Armed Forces Policy Council; Army Ground Forces -> Chief of Army Field Forces;
  Division of Territories -> Office of Territories; Public Buildings Administration -> Public Buildings
  Service; Procurement Division -> Federal Supply Service; Bureau of Community Facilities -> Community
  Facilities Service (moved to HHFA); Office of Domestic Commerce -> Office of Industry and Commerce;
  Office of Research and Intelligence -> Office of Intelligence Research; Office of Intelligence
  Collection and Dissemination -> Office of Libraries and Intelligence Acquisition; State economic
  "Divisions" -> "Staffs"; Post Office "Office of the Assistant Solicitor General" etc.
Public Roads Administration -> Bureau of Public Roads: the sheet already has that row (43), so 430 Dead
and 43 is unhidden (FAQ).
"""
import openpyxl, shutil

SRC = 'data/Spreadsheet 1935-1949.xlsx'
DST = 'data/Spreadsheet 1935-1950.xlsx'
YEAR, PREV_YEAR = 1950, 1949

AG, COM, JUS, LAB, ST = ('Department of Agriculture', 'Department of Commerce', 'Department of Justice',
                         'Department of Labor', 'Department of State')
INT, TRE, PO = 'Department of the Interior', 'Department of the Treasury', 'Post Office Department'
NME, ARMY = 'National Military Establishment', 'Department of the Army'
CABINET = {AG, COM, JUS, LAB, ST, INT, TRE, PO, NME}
DEAD, NOTYET, OTHER, CAB = 'DEAD', 'NOTYET', 'OTHER', 'CAB'
GSA, HHFA, FSA, BES = 'General Services Administration', 'Housing and Home Finance Agency', 'Federal Security Agency', 'Bureau of Employment Security'
NARS = 'National Archives and Records Service'
BPOO, BFAC, BFIN, BTR = 'Bureau of Post Office Operations', 'Bureau of Facilities (Post Office)', 'Bureau of Finance (Post Office)', 'Bureau of Transportation (Post Office)'
EURB, ARAB, FEB = 'Bureau of European Affairs', 'Bureau of Inter-American Affairs', 'Bureau of Far Eastern Affairs'
NEAB, UNAB, GERB = 'Bureau of Near Eastern, South Asian, and African Affairs', 'Bureau of United Nations Affairs', 'Bureau of German Affairs'
IAC, FDP, TCP, ITP = 'Office of International Administration and Conferences', 'Office of Financial and Development Policy', 'Office of Transportation and Communications', 'Office of International Trade Policy'


def vals(c):
    """Five-column coding: cabinet, child, cabinet parent, other parent, other agency."""
    if c == DEAD: return ['Dead'] * 5
    if c == NOTYET: return ['No', 'No', 'None', 'None', 'No']
    if c == CAB: return ['Yes', 'No', 'None', 'None', 'No']
    if c == OTHER: return ['No', 'No', 'None', 'None', 'Yes']
    if c in CABINET: return ['No', 'Yes', c, 'None', 'No']
    return ['No', 'Yes', 'None', c, 'No']


O = {}
NAMED = []
def rows(code, *row_numbers):
    O.update({r: code for r in row_numbers})
def named(code, *names):
    NAMED.extend((n, code) for n in names)

# Executive Office (chart p. 493) and independent agencies
rows(DEAD, 682, 747)                             # National Intelligence Authority (appendix only); Advisory Committee on Voluntary Foreign Aid (not printed)
rows(OTHER, 8, 457)                              # Arlington Memorial Amphitheater Commission, Panama Canal: own Independent Agencies headings (pp. 287, 396)
# Department of Defense (chart p. 500); Navy chart p. 502 unchanged
rows(DEAD, 811)                                  # Aeronautical Board: appendix run-in only
# Army (chart p. 501): no Special Staff box; its units hang directly on the Army line
rows(DEAD, 270, 776, 780, 781)                   # Special Staff, Troop Information and Education, Civil Affairs, Special Services: not on chart
rows(ARMY, 275, 283, 294, 539, 775, 777)         # Inspector General, Information (Public Relations), National Guard Bureau, Reserve and ROTC, Legislative Liaison, Military History
# State (chart p. 498; officials pp. 68-80): regional Divisions replaced by Offices under Bureaus
rows(DEAD, 572, 573, 574, 575, 825, 860)         # Offices of American Republic, European, Near Eastern, Far Eastern, UN, German-Austrian Affairs -> Bureaus
rows(DEAD, 127, 142, 143, 149, 589, 591, 594, 595, 599, 600, 603, 660, 661, 663, 664, 755, 756, 757, 758,
     828, 829, 830, 831, 864, 865, 866, 867, 868, 869)   # regional/UN/German Divisions -> new Offices (rows below)
rows(DEAD, 666, 832, 740)                        # Labor-Social division, Economic-Property division, Map Intelligence: not listed
rows(IAC, 140)                                   # Division of International Conferences
# Treasury (chart p. 499, list p. 86)
# Interior (chart p. 506)
# Commerce (chart p. 508): Maritime Commission abolished May 24, 1950 -> Maritime Administration, Federal Maritime Board
rows(DEAD, 469, 430)
rows(COM, 43)                                    # Bureau of Public Roads transferred to Commerce (Reorg. Plan 7 of 1949) - FAQ: unhide
# Labor (chart p. 509): Bureau of Employment Security returned to Labor Aug 20, 1949; Employees' Compensation May 24, 1950
rows(LAB, 484, 795, 796)
# Federal Security Agency (chart p. 516)
rows(DEAD, 792, 798, 800)                        # Appeals Council, Office of Research, Office of Information: not on chart
# General Services Administration (chart p. 519; section p. 353)
rows(DEAD, 332)                                  # Federal Real Estate Board: index only
rows(HHFA, 652)                                  # Community Facilities Service (chart p. 521)
rows(NARS, 360, 361, 366, 345)                   # National Archives, its Council and NHPC ("Related Organizations" p. 361-362), FDR Library
# Housing and Home Finance Agency (chart p. 521)
rows(DEAD, 323)                                  # Federal Home Loan Bank System: sentences only
# Post Office (chart p. 505; list p. 170): Assistant PMG bureaus abolished Aug 20, 1949 (Reorg. Plan 3)
rows(DEAD, 388, 394, 401, 411, 409, 418, 643, 558)
rows(BPOO, 870, 391, 872, 874, 873, 871)
rows(BFAC, 412, 413, 414, 415, 416, 417)
rows(BFIN, 407, 642, 402, 406, 404, 641, 403)
rows(BTR, 773, 396, 772, 398)
# Agriculture (chart p. 507)
rows(DEAD, 841)                                  # Administrator, Research and Marketing Act: not on chart

NEW = [
    ('Air Coordinating Committee', OTHER),                          # p. 286
    ('Commodity Exchange Commission', OTHER),                       # p. 295
    ('National Advisory Council on International Monetary and Financial Problems', OTHER),  # p. 384
    ('National Forest Reservation Commission', OTHER),              # p. 388
    ('National Science Foundation', OTHER),                         # p. 394
    ('Panama Railroad Company', 'The Panama Canal'),                # p. 397
    ('Fair Employment Board', 'United States Civil Service Commission'),   # p. 435
    # Department of Defense (chart p. 500)
    ('Office of Medical Services (Defense)', NME), ('Office of Public Information (Defense)', NME),
    ('Staff Council (Defense)', NME), ('Management Committee (Defense)', NME),
    ('Comptroller of the Army', ARMY),                              # chart p. 501
    # State
    ('Mutual Defense Assistance Program', ST),                      # chart p. 498, officials p. 68
    (EURB, ST), (ARAB, ST), (FEB, ST), (NEAB, ST), (UNAB, ST), (GERB, ST),
    ('Office of Eastern European Affairs', EURB), ('Office of Western European Affairs', EURB),
    ('Office of British Commonwealth and Northern European Affairs', EURB), ('Office of European Regional Affairs', EURB),
    ('Office of East Coast Affairs', ARAB), ('Office of North and West Coast Affairs', ARAB),
    ('Office of Middle American Affairs', ARAB), ('Office of Regional American Affairs', ARAB),
    ('Office of Chinese Affairs', FEB), ('Office of Northeast Asian Affairs', FEB), ('Office of Philippine and Southeast Asian Affairs', FEB),
    ('Office of Greek, Turkish, and Iranian Affairs', NEAB), ('Office of African and Near Eastern Affairs', NEAB), ('Office of South Asian Affairs', NEAB),
    ('Office of United Nations Economic and Social Affairs', UNAB), ('Office of Dependent Area Affairs', UNAB),
    (IAC, UNAB), ('Office of United Nations Political and Security Affairs', UNAB), ('Division of International Administration', IAC),
    ('Office of German Political Affairs', GERB), ('Office of German Economic Affairs', GERB), ('Office of German Public Affairs', GERB),
    ('Inland Transport Policy Staff', TCP), ('Shipping Policy Staff', TCP),
    ('International Business Practices Policy Staff', ITP),
    ('Investment and Economic Development Staff', FDP), ('Lend-Lease and Surplus Property Staff', FDP),
    # Treasury (chart p. 499)
    ('Office of Budget (Treasury)', TRE),
    # Interior (chart p. 506)
    ('Southeastern Power Administration', INT), ('Division of Management Research (Interior)', INT), ('Division of Property Management (Interior)', INT),
    # Commerce (chart p. 508)
    ('Maritime Administration', COM), ('Federal Maritime Board', COM),
    ('Office of Program Planning (Commerce)', COM), ('Office of Publications (Commerce)', COM),
    # Labor (chart p. 509)
    ("Veterans' Employment Service", LAB),
    # Federal Security Agency (chart p. 516)
    ('Office of Field Services (FSA)', FSA),
    # General Services Administration (chart p. 519)
    (NARS, GSA), ('Federal Register Division', NARS), ('Contract Settlement Appeal Board', GSA),
    ('Office of Comptroller (GSA)', GSA), ('Office of Management (GSA)', GSA), ('Office of General Counsel (GSA)', GSA),
    ('Office of Public Information and Reports (GSA)', GSA), ('Compliance Division (GSA)', GSA),
    # Housing and Home Finance Agency (chart p. 521)
    ('Division of Slum Clearance and Urban Redevelopment', HHFA), ('Division of Housing Research', HHFA), ('Division of College Housing', HHFA),
    # Post Office (chart p. 505)
    (BPOO, PO), (BFAC, PO), (BFIN, PO), (BTR, PO),
    ('Division of Equipment and Supplies (Post Office)', BFAC),
    ('Division of Highway and Contract Transportation', BTR), ('Division of Transportation Accounts', BTR),
]

shutil.copy(SRC, DST)
wb = openpyxl.load_workbook(DST)
ws = wb.worksheets[0]
hdr = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value}
COL, PREV = hdr[f'{YEAR}_cabinet_agency'], hdr[f'{PREV_YEAR}_cabinet_agency']

last = max(r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value)
byname = {}
for r in range(3, last + 1):
    if ws.cell(r, 1).value:
        byname.setdefault(str(ws.cell(r, 1).value).strip(), []).append(r)
for n, c in NAMED:
    assert n in byname, n
    for r in byname[n]:
        O[r] = c
for r in range(3, last + 1):
    if not ws.cell(r, 1).value:
        continue
    prev = [ws.cell(r, PREV + i).value for i in range(5)]
    if r in O:
        new = vals(O[r])
        if O[r] not in (DEAD, NOTYET) and not (ws.cell(r, 2).value or '').strip():
            ws.cell(r, 2).value = 'Executive'
    elif all(v in (None, '') for v in prev):
        continue
    else:
        new = prev
    for i, v in enumerate(new):
        ws.cell(r, COL + i).value = v

r = last + 1
for name, c in NEW:
    ws.cell(r, 1).value = name
    ws.cell(r, 2).value = 'Executive'
    for col in range(3, COL):
        ws.cell(r, col).value = vals(NOTYET)[(col - 3) % 5]
    for i, v in enumerate(vals(c)):
        ws.cell(r, COL + i).value = v
    r += 1

# FAQ: hide (never delete) dead rows and legislative/judicial rows
for x in range(3, ws.max_row + 1):
    if ws.cell(x, 1).value:
        kind = (ws.cell(x, 2).value or '').strip()
        ws.row_dimensions[x].hidden = kind in ('Legislative', 'Judicial') or ws.cell(x, COL).value == 'Dead'

wb.save(DST)
print('saved', DST, 'new rows', last + 1, '-', r - 1)
