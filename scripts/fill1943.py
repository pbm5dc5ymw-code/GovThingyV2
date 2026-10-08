"""Fill the 1943 columns from the Winter 1943-44 Government Manual (December 1, 1943).

That edition is the most recent of 1943 (the FAQ rule). It has no organization
charts, so structure comes from page images of each department's officials page
and indented unit list, and from the Executive Office / OEM membership lists
(printed pages 53 and 61). Each row carries its 1942 coding unless changed below;
page numbers are the printed page numbers in that manual.
"""
import openpyxl, shutil

SRC = 'data/Spreadsheet 1935-1942 (Sept 1942 manual).xlsx'
DST = 'data/Spreadsheet 1935-1943.xlsx'
COL, PREV = 38, 33  # 1943_cabinet_agency, 1942_cabinet_agency

AG, COM, JUS, ST = 'Department of Agriculture', 'Department of Commerce', 'Department of Justice', 'Department of State'
INT, NAVY, TRE, WAR, PO = ('Department of the Interior', 'Department of the Navy', 'Department of the Treasury',
                           'Department of War', 'Post Office Department')
CABINET = {AG, COM, JUS, ST, INT, NAVY, TRE, WAR, PO, 'Department of Labor'}
EOP, OEM, FSA = 'Executive Office of the President', 'Office for Emergency Management', 'Federal Security Agency'
WMC, WPB, FEA = 'War Manpower Commission', 'War Production Board', 'Foreign Economic Administration'
WFA, FPA, ASF = 'War Food Administration', 'Food Production Administration', 'Army Service Forces'
B4 = 'Bureau of the Fourth Assistant Postmaster General'
DEAD, OTHER = 'DEAD', 'OTHER'


def vals(c):
    if c == DEAD: return ['Dead'] * 5
    if c == OTHER: return ['No', 'No', 'None', 'None', 'Yes']
    if c in CABINET: return ['No', 'Yes', c, 'None', 'No']
    return ['No', 'Yes', 'None', c, 'No']


O = {}
def rows(code, *rs):
    O.update({r: code for r in rs})

# Executive Office (p. 53) and OEM membership (p. 61)
rows(EOP, 378)                                  # Committee for Congested Production Areas
rows(DEAD, 374)                                 # NRPB abolished Aug 31, 1943 (p. 57)
rows(OEM, 510, 495, 380)                        # Office of Economic Stabilization, Office of War Mobilization, FEA
rows(DEAD, 493, 512, 525, 491)                  # Lend-Lease, BEW/OEW and Export Control -> FEA; ODHWS abolished Apr 1943
rows(FEA, 304)                                  # Export-Import Bank transferred to FEA (p. 68)
rows(OTHER, 506)                                # Board of War Communications: own section, not in OEM list (p. 62)
rows(WMC, 516, 105, 311, 545)                   # Bureau of Selective Service, USES, apprenticeship, Training Within Industry (pp. 100-104)
rows(DEAD, 375)                                 # NYA liquidated 1943
rows(OTHER, 206)                                # Petroleum Administration for War (p. 156)
rows('Joint Chiefs of Staff', 542)              # OSS under the Joint Chiefs of Staff (p. 161)
# Agriculture (pp. 328-329): War Food Administration
rows(WFA, 24)                                   # Commodity Credit Corporation
rows(FPA, 4, 306, 45, 54)                       # AAA, Farm Security, FCIC, SCS within Food Production Administration
rows(DEAD, 30, 540, 59, 55)                     # AMA, ACAA, Agricultural War Relations, Sugar Agency consolidated
# Interior (p. 306)
rows(INT, 205, 179, 204)                        # Solid Fuels Administration for War, Coal Mines Administration, Southwestern Power Administration
rows(DEAD, 200, 541, 363, 208)                  # replaced coordinators, Bituminous Coal Division, War Resources Council
# Navy (p. 275)
rows(DEAD, 544)                                 # Office of Procurement and Material not listed
# Treasury (p. 206)
rows(TRE, 249, 263)                             # War Finance Division, Office of the Tax Legislative Counsel
rows(DEAD, 244, 262)                            # War Savings Staff, Processing Tax Board of Review
# War (pp. 235-236): Services of Supply renamed Army Service Forces
rows(WAR, 272)
rows(DEAD, 271)
rows(ASF, 274, 294, 539, 277, 278, 276, 536, 291, 293, 296, 279, 280, 292, 281)
# Justice (p. 250)
rows(JUS, 95)                                   # Office of the Attorney General
rows(DEAD, 77, 79)                              # Bond and Spirits Division, Bureau of War Risk Litigation
# Post Office (p. 261)
rows(PO, 427)                                   # Office of the Postmaster General
rows(DEAD, 391, 392, 399)                       # Divisions of Postmasters, Dead Letters, Rural Mails
rows(B4, 418)                                   # Division of Mail Equipment Shops
# State (p. 182)
rows(ST, 145)                                   # Division of Political Studies
rows(DEAD, 543, 130, 128, 121, 136, 151, 152, 112)
# Federal Works (p. 445) and Commerce
rows(DEAD, 317, 477, 344)                       # PWA and WPA functions in liquidation in the Administrator's office
rows(DEAD, 300)                                 # Electric Home and Farm Authority terminated Oct 1942
rows(DEAD, 356, 442, 454)                       # Maritime Labor Board (expired June 1942), Savings and Loan Division, Codification Board: not in 1943 text

NEW = [
    ('Committee on Fair Employment Practice', OEM),
    ('Smaller War Plants Corporation', WPB),
    ('Joint Chiefs of Staff', OTHER),
    ('War Food Administration', AG),
    ('Food Distribution Administration', WFA),
    ('Food Production Administration', WFA),
    ('Budget and Administrative Management Division', INT),
    ('Division of Personnel Supervision and Management (Interior)', INT),
    ('Division of Classification (Interior)', INT),
    ('Executive Office of the Secretary (Navy)', NAVY),
    ('Board of Immigration Appeals', JUS),
    ('Office of Budget and Administrative Planning', PO),
    ('Division of Economic Studies', ST),
    ('Office of Community War Services', FSA),
    ('Office of Vocational Rehabilitation', FSA),
]

shutil.copy(SRC, DST)
wb = openpyxl.load_workbook(DST)
ws = wb.worksheets[0]
assert ws.cell(2, COL).value == '1943_cabinet_agency'

last = max(r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value)
for r in range(3, last + 1):
    if not ws.cell(r, 1).value:
        continue
    prev = [ws.cell(r, PREV + i).value for i in range(5)]
    if r in O:
        new = vals(O[r])
        if O[r] != DEAD and not (ws.cell(r, 2).value or '').strip():
            ws.cell(r, 2).value = 'Executive'
    elif all(v in (None, '') for v in prev):
        continue
    else:
        new = prev
    for i, v in enumerate(new):
        ws.cell(r, COL + i).value = v

r = last + 1
for name, c in NEW:
    ws.cell(r, 1).value = name
    ws.cell(r, 2).value = 'Executive'
    for col in range(3, COL):
        ws.cell(r, col).value = ['No', 'No', 'None', 'None', 'No'][(col - 3) % 5]
    for i, v in enumerate(vals(c)):
        ws.cell(r, COL + i).value = v
    r += 1

# FAQ: hide (never delete) dead and legislative/judicial rows
for x in range(3, ws.max_row + 1):
    if ws.cell(x, 1).value:
        kind = (ws.cell(x, 2).value or '').strip()
        ws.row_dimensions[x].hidden = kind in ('Legislative', 'Judicial') or ws.cell(x, COL).value == 'Dead'

wb.save(DST)
print('saved', DST, 'new rows', last + 1, '-', r - 1)
