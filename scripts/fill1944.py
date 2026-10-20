"""Fill the 1944 columns from the Summer 1944 Government Manual (June 1, 1944; text current to about August 1944).

That is the only 1944 edition on govinfo. It has no organization charts, so structure comes from page
images of each department's officials page and "as follows" unit list, the Executive Office and OEM
membership lists (printed pp. 53, 61), and Appendix A. Each row carries its 1943 coding unless changed
below; page numbers are printed page numbers in that manual (PDF page = printed + 8).
"""
import openpyxl, shutil

SRC = 'data/Spreadsheet 1935-1943.xlsx'
DST = 'data/Spreadsheet 1935-1944.xlsx'
YEAR, PREV_YEAR = 1944, 1943

AG, COM, JUS, LAB, ST = ('Department of Agriculture', 'Department of Commerce', 'Department of Justice',
                         'Department of Labor', 'Department of State')
INT, NAVY, TRE, WAR, PO = ('Department of the Interior', 'Department of the Navy', 'Department of the Treasury',
                           'Department of War', 'Post Office Department')
CABINET = {AG, COM, JUS, LAB, ST, INT, NAVY, TRE, WAR, PO}
DEAD, NOTYET, OTHER, CAB = 'DEAD', 'NOTYET', 'OTHER', 'CAB'
EOP, OWM, FEA = 'Executive Office of the President', 'Office of War Mobilization', 'Foreign Economic Administration'
WFA, ARA, ASF, JCS = 'War Food Administration', 'Agricultural Research Administration', 'Army Service Forces', 'Joint Chiefs of Staff'
B3 = 'Bureau of the Third Assistant Postmaster General'

# State Department offices after the January 15, 1944 reorganization (pp. 185-188)
SPA, ARA_ST, EUR = 'Office of Special Political Affairs', 'Office of American Republic Affairs', 'Office of European Affairs'
NEA, FE, CTL = 'Office of Near Eastern and African Affairs', 'Office of Far Eastern Affairs', 'Office of Controls'
TC, WEA, ECA = 'Office of Transportation and Communications', 'Office of Wartime Economic Affairs', 'Office of Economic Affairs'
PI, DA, FS = 'Office of Public Information', 'Office of Departmental Administration', 'Office of the Foreign Service'


def vals(c):
    """Five-column coding: cabinet, child, cabinet parent, other parent, other agency."""
    if c == DEAD: return ['Dead'] * 5
    if c == NOTYET: return ['No', 'No', 'None', 'None', 'No']
    if c == CAB: return ['Yes', 'No', 'None', 'None', 'No']
    if c == OTHER: return ['No', 'No', 'None', 'None', 'Yes']
    if c in CABINET: return ['No', 'Yes', c, 'None', 'No']
    return ['No', 'Yes', 'None', c, 'No']


O = {}
def rows(code, *row_numbers):
    O.update({r: code for r in row_numbers})

# Executive Office (p. 53) and war agencies
rows(INT, 499)                                  # War Relocation Authority transferred to Interior Feb 16, 1944 (p. 318 list; App. A p. 639)
rows(OTHER, 516)                                # Selective Service System made a separate agency Dec 5, 1943 (p. 156)
rows(OTHER, 511, 514)                           # OPA and Office of Censorship: own sections, not in the EOP list (p. 53) or OEM list (p. 61); no charts, so the 1942 chart footnote no longer applies
# State: Departmental Order 1218 of Jan 15, 1944 abolished these offices and divisions (App. A pp. 635-636);
# the rest are not in the new officials list (pp. 185-187)
rows(DEAD, 111, 116, 118, 119, 120, 125, 126, 132, 133, 134, 139, 145, 155, 156, 157, 158, 160, 161,
     167, 168, 169, 533, 559)
rows(DEAD, 113, 114, 131, 153, 165)             # boards, committee, training school, liaison office: no longer units (pp. 185-187, 201-202)
rows(DEAD, 23, 342)                             # Committee on Trade Agreements, Foreign Service Buildings Commission: no section
rows(DA, 123, 129, 140, 146)                    # Communications and Records, Departmental Personnel, International Conferences, Protocol (pp. 186-187)
rows(CTL, 135, 166, 170)                        # Foreign Activity Correlation, Passport, Visa (p. 186)
rows(FS, 137, 138)                              # Foreign Service Administration, Foreign Service Personnel (p. 187)
rows(NEA, 143)                                  # Division of Near Eastern Affairs (p. 186)
rows(PI, 147)                                   # Division of Research and Publication (p. 186)
rows(WEA, 150)                                  # World Trade Intelligence Division (p. 186)
rows(ARA_ST, 142)                               # Division of Mexican Affairs reappears (p. 185) - FAQ: unhide and relabel
rows(EUR, 127, 149)                             # Divisions of Eastern and Western European Affairs reappear (p. 186)
# Agriculture / War Food Administration (pp. 342-344, 364-385)
rows(WFA, 4, 306, 54)                           # Agricultural Adjustment Agency, Farm Security Administration, Soil Conservation Service report to the War Food Administrator (p. 342-343)
# Interior (p. 318)
rows(DEAD, 555)                                 # Division of Classification not in the unit list
# Post Office (p. 272)
rows(DEAD, 405)                                 # Division of Classification replaced by two mail divisions
# Federal Security Agency (p. 437)
rows(DEAD, 20, 517)                             # CCC liquidated by June 30, 1943 (App. A); not a constituent unit

# --- Changes after comparing with the hand-entered 1944 data and re-checking page layout ---
# Mentioned only in a sentence of another unit's paragraph: no heading, list entry or officials block
rows(DEAD, 331)                                 # Federal Prison Industries: one sentence under Bureau of Prisons (p. 267); not in Justice list (p. 260)
rows(DEAD, 437)                                 # Defense Homes Corporation: one sentence in FPHA creation text (p. 139)
rows(DEAD, 314)                                 # "Federal credit unions": a run-in activity paragraph in the FDIC section (p. 485)
rows(DEAD, 3)                                   # Advisory Council of the National Arboretum: one sentence under Bureau of Plant Industry (p. 347)
# Duplicate rows: the unit list names the Bureau, and the hand sheet keeps the Bureau row
rows(DEAD, 389, 393, 400, 410, 419)             # "Office of the ... Assistant Postmaster General", "Chief Post Office Inspector" (list p. 272 has Bureaus)
rows(DEAD, 102)                                 # second Immigration and Naturalization Service row (row 85 is live)
rows(JUS, 96)                                   # Administrative Division: hand sheet uses this row; same unit as row 74 (p. 260)
rows('Farm Credit Administration', 302)         # Cooperative Research and Service Division: centered subheading in FCA section (p. 357), same format as RACC
rows(OTHER, 361, 366)                           # National Archives Council, National Historical Publications Commission: own top-level headings (pp. 527-528)

NEW = [
    ('War Refugee Board', EOP),                                     # p. 59
    ('Surplus War Property Administration', OWM),                   # p. 98-99
    ('Retraining and Reemployment Administration', OWM),            # p. 98-99
    ('Office of Contract Settlement', OTHER),                       # p. 141-142
    ('U. S. Commercial Company', FEA),                              # p. 67, footnote p. 67
    ('Petroleum Reserves Corporation', FEA),                        # p. 66
    ('Rubber Development Corporation', FEA),                        # p. 66
    ('Committee on Medical Research', 'Office of Scientific Research and Development'),   # p. 88
    ('Army and Navy Staff College', JCS),                           # p. 162
    # State offices (p. 188) and their divisions (pp. 185-187)
    (SPA, ST), (ARA_ST, ST), (EUR, ST), (NEA, ST), (FE, ST), (CTL, ST),
    (TC, ST), (WEA, ST), (ECA, ST), (PI, ST), (DA, ST), (FS, ST),
    ('Policy Committee', ST),                                       # p. 185, 189
    ('Committee on Post-War Programs', ST),                         # p. 185, 190
    ('Executive Committee on Economic Foreign Policy', ST),         # p. 205
    ('Division of International Security and Organization', SPA),
    ('Division of Territorial Studies', SPA),
    ('Division of Caribbean and Central American Affairs', ARA_ST),
    ('Division of Brazilian Affairs', ARA_ST),
    ('Division of Bolivarian Affairs', ARA_ST),
    ('Division of River Plate Affairs', ARA_ST),
    ('Division of West Coast Affairs', ARA_ST),
    ('Division of American Republics Analysis and Liaison', ARA_ST),
    ('Division of British Commonwealth Affairs', EUR),
    ('Division of Northern European Affairs', EUR),
    ('Division of Southern European Affairs', EUR),
    ('Division of Central European Affairs', EUR),
    ('Division of Middle Eastern Affairs', NEA),
    ('Division of African Affairs', NEA),
    ('Division of Chinese Affairs', FE),
    ('Division of Japanese Affairs', FE),
    ('Division of Southwest Pacific Affairs', FE),
    ('Division of Philippine Affairs', FE),
    ('Special War Problems Division', CTL),
    ('Aviation Division', TC),
    ('Shipping Division', TC),
    ('Telecommunications Division', TC),
    ('Supply and Resources Division', WEA),
    ('Liberated Areas Division', WEA),
    ('American Republics Requirements Division', WEA),
    ('Eastern Hemisphere Division', WEA),
    ('Division of Commercial Policy', ECA),
    ('Division of Financial and Monetary Affairs', ECA),
    ('Commodities Division', ECA),
    ('Petroleum Division', ECA),
    ('Division of Labor Relations', ECA),
    ('Division of Public Liaison', PI),
    ('Motion Picture and Radio Division', PI),
    ('Science, Education, and Art Division', PI),
    ('Central Translating Division', PI),
    ('Division of Coordination and Review', DA),
    ('Division of Budget and Finance', DA),
    ('Division of Administrative Management', DA),
    ('Division of Geography and Cartography', DA),
    ('Division of Foreign Buildings Operations', FS),
    # Agriculture / War Food Administration
    ('Extension Service', WFA),                                     # p. 342, 374-375
    ('Office of Labor (WFA)', WFA),                                 # p. 343, 380
    ('Office of Materials and Facilities (WFA)', WFA),              # p. 343, 381
    ('Office of Price (WFA)', WFA),                                 # p. 343, 381
    ('Office of Transportation (WFA)', WFA),                        # p. 343, 385
    ('Beltsville Research Center', ARA),                            # p. 342, 352
    ('Office of Budget and Finance (Agriculture)', AG),             # staff offices p. 343-344, 386-389
    ('Office of Information (Agriculture)', AG),
    ('Office of Personnel (Agriculture)', AG),
    ('Office of Plant and Operations (Agriculture)', AG),
    ('Library (Agriculture)', AG),
    ('Office of the Solicitor (Agriculture)', AG),
    # War, Justice, Post Office
    ('Office of the Chief of Transportation', ASF),                 # p. 245
    ('Office of the Solicitor General', JUS),                       # p. 260
    ('Office of the Chief Clerk and Director of Personnel', PO),    # p. 272
    ('Division of Letter and Miscellaneous Mails', B3),             # p. 272
    ('Division of Newspaper and Periodical Mail', B3),              # p. 272
    ('Division of Methods and Procedures', 'Bureau of Accounts'),   # p. 272 (Post Office Bureau of Accounts)
    # Federal Security Agency and independent agencies
    ('Bureau of Old-Age and Survivors Insurance', 'Social Security Board'),   # p. 451-452
    ('National Capital Housing Authority', OTHER),                  # p. 529
]

shutil.copy(SRC, DST)
wb = openpyxl.load_workbook(DST)
ws = wb.worksheets[0]
hdr = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value}
COL, PREV = hdr[f'{YEAR}_cabinet_agency'], hdr[f'{PREV_YEAR}_cabinet_agency']

last = max(r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value)
for r in range(3, last + 1):
    if not ws.cell(r, 1).value:
        continue
    prev = [ws.cell(r, PREV + i).value for i in range(5)]
    if r in O:
        new = vals(O[r])
        if O[r] not in (DEAD, NOTYET) and not (ws.cell(r, 2).value or '').strip():
            ws.cell(r, 2).value = 'Executive'
    elif all(v in (None, '') for v in prev):
        continue
    else:
        new = prev
    for i, v in enumerate(new):
        ws.cell(r, COL + i).value = v

r = last + 1
for name, c in NEW:
    ws.cell(r, 1).value = name
    ws.cell(r, 2).value = 'Executive'
    for col in range(3, COL):
        ws.cell(r, col).value = vals(NOTYET)[(col - 3) % 5]
    for i, v in enumerate(vals(c)):
        ws.cell(r, COL + i).value = v
    r += 1

# FAQ: hide (never delete) dead rows and legislative/judicial rows
for x in range(3, ws.max_row + 1):
    if ws.cell(x, 1).value:
        kind = (ws.cell(x, 2).value or '').strip()
        ws.row_dimensions[x].hidden = kind in ('Legislative', 'Judicial') or ws.cell(x, COL).value == 'Dead'

wb.save(DST)
print('saved', DST, 'new rows', last + 1, '-', r - 1)
