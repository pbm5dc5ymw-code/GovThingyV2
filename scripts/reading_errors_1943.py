"""Build a table of 1943 reading errors on both sides (my sheet vs. hand-entered sheet).

Pages are printed page numbers in the Winter 1943-44 Government Manual (Dec. 1, 1943).
"Wrong in June 1943 edition too?" says whether the entry would also be wrong if the
June 1943 manual had been the source.
"""
import openpyxl, re
from openpyxl.styles import Font, Alignment, PatternFill

MINE = 'data/Spreadsheet 1935-1943.xlsx'
HAND = '/Users/leo/Downloads/_FeedBack Spreadsheet_ (1).xlsx'
OUT = 'data/1943 Reading Errors.xlsx'
COL = 38


def code(ws, r):
    v = [str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else '' for c in range(COL, COL + 5)]
    if not any(v): return '(blank)'
    if v[0] == 'Dead': return 'Dead'
    if v[0] == 'Yes': return 'Cabinet agency'
    if v[1] == 'Yes': return 'Child of ' + (v[2] if v[2] not in ('None', '') else v[3])
    if v[4] == 'Yes': return 'Other agency'
    return 'Not listed (No/No/None/None/No)'


mine_ws = openpyxl.load_workbook(MINE).worksheets[0]
hand_ws = openpyxl.load_workbook(HAND)['Agencies and Years']
hand_rows = {re.sub(r'[^a-z]', '', str(hand_ws.cell(r, 1).value or '').lower()): r for r in range(3, hand_ws.max_row + 1)}

# (side, row, my coding override or None, what the manual shows, pages, explanation, wrong in June too)
E = [
    # --- My errors ---
    ('Mine', 144, 'Child of Department of State', 'Not in the list of State divisions and offices', '182',
     'Carried forward from 1942 without checking the 1943 list. Fixed.', 'Yes'),
    ('Mine', 231, 'Child of Department of the Navy', 'Not in the list of Navy principal divisions; the Judge Advocate General is named only as an adviser', '275',
     'Carried forward from 1942; an official was treated as a unit. Fixed.', 'Yes'),
    ('Mine', 322, 'Child of Federal Home Loan Bank Administration', 'The Bank Administration "performs functions formerly exercised by the Federal Home Loan Bank Board"', '137',
     'Board no longer exists as a unit; carried forward. Fixed.', 'Yes'),
    ('Mine', 159, 'Child of Department of State', '"Adviser on International Economic Affairs (Vacancy)" appears only in the officials list', '180',
     'An official post, not an office; carried forward. Fixed.', 'Yes'),
    ('Mine', 484, 'Child of Federal Security Agency', 'Bureaus are printed inside the Social Security Board section', '433-435',
     'Wrong parent level (also in 1942). Caught by my own audit and fixed before the comparison.', 'Yes'),
    ('Mine', 323, 'Child of National Housing Agency', 'Bank Administration supervises the Federal Home Loan Bank System', '137',
     'Wrong parent level (also in 1942; same for HOLC and FSLIC). Caught by my own audit and fixed before the comparison.', 'Yes'),
    ('Mine', 356, 'Other agency (1942)', 'Maritime Labor Board title expired June 22, 1942 (Appendix A)', '615',
     'Kept alive in 1942. Caught by my own audit and fixed before the comparison.', 'Yes'),
    # --- Hand-entered: agency exists but marked Dead or Not listed ---
    ('Hand', 22, None, 'Has its own section under State related organizations', '195', 'Agency still exists; marked Dead.', 'Yes'),
    ('Hand', 116, None, 'In the list of State divisions and offices', '182', 'Still listed; marked Dead.', 'Yes'),
    ('Hand', 160, None, 'In the list of State divisions and offices', '182', 'Still listed; marked Dead.', 'Yes'),
    ('Hand', 227, None, 'In the list of Navy principal divisions', '275', 'Still listed; marked Dead.', 'Yes'),
    ('Hand', 258, None, 'In the list of Treasury principal branches', '206', 'Still listed; marked Dead.', 'Yes'),
    ('Hand', 408, None, 'Listed under the Bureau of Accounts', '261', 'Still listed; marked Dead.', 'Yes'),
    ('Hand', 323, None, 'Supervised by the Federal Home Loan Bank Administration', '137', 'Still exists; marked Dead.', 'Yes'),
    ('Hand', 349, None, 'Own heading in the National Housing Agency section', '136, 139', 'Still exists; marked Dead.', 'Yes'),
    ('Hand', 437, None, 'Administered by the Federal Public Housing Authority', '146', 'Still exists; marked Dead.', 'Yes'),
    ('Hand', 438, None, 'Own section under the Farm Credit Administration', '342-343', 'Still exists; marked Dead.', 'Yes'),
    ('Hand', 331, None, 'Operates under supervision of the Bureau of Prisons', '257', 'Still exists; marked Dead.', 'Yes'),
    ('Hand', 462, None, 'Own section (members are six Secretaries)', '60', 'Still exists; marked Dead.', 'Yes'),
    ('Hand', 269, None, 'War Department General Staff in the officials list', '234', 'Exists; coded as not listed.', 'Yes'),
    ('Hand', 272, None, 'Own heading in the War Department officials (created March 9, 1943)', '235', 'Exists; coded as not listed.', 'Yes'),
    ('Hand', 273, None, 'Own heading in the War Department officials', '236', 'Exists; coded as not listed.', 'Yes'),
    ('Hand', 286, None, 'Own heading in the War Department officials', '235', 'Exists; coded as not listed.', 'Yes'),
    ('Hand', 283, None, 'Bureau of Public Relations in the War Department officials', '234', 'Exists; coded as not listed.', 'Yes'),
    ('Hand', 457, None, 'Own section in the War Department chapter', '237', 'Exists; coded as not listed.', 'Yes'),
    ('Hand', 8, None, 'Own section in the War Department chapter', '237', 'Exists; coded as not listed.', 'Yes'),
    # --- Hand-entered: agency gone but marked alive ---
    ('Hand', 376, None, 'Consolidated into the Office of War Information, June 13, 1942', '53, 92', 'Abolished in 1942; still coded alive.', 'Yes'),
    ('Hand', 112, None, 'Not in the list of State divisions and offices (also absent in 1942)', '182', 'No longer listed; coded alive.', 'Yes'),
    ('Hand', 122, None, 'Replaced in 1942 by the Division of Commercial Policy and Agreements', '182', 'Renamed unit; old name coded alive.', 'Yes'),
    ('Hand', 30, None, 'Consolidated into Food Distribution Administration, Dec. 5, 1942', '598', 'Abolished before 1943; coded alive.', 'Yes'),
    ('Hand', 50, None, 'Became a branch of the Agricultural Marketing Administration, Feb. 1942', '598', 'Abolished before 1943; coded alive.', 'Yes'),
    ('Hand', 59, None, 'Functions moved to Food Production/Distribution Administrations, Dec. 5, 1942', '618', 'Abolished before 1943; coded alive.', 'Yes'),
    ('Hand', 504, None, 'Renamed Office for Agricultural War Relations, Jan. 1942', '618', 'Renamed before 1943; coded alive.', 'Yes'),
    ('Hand', 491, None, 'Abolished April 29, 1943; functions to Federal Security Agency', '605; 421', 'Abolished before June 1943; coded alive.', 'Yes'),
    ('Hand', 505, None, 'Renamed Board of Economic Warfare, Dec. 1941', '598', 'Renamed before 1943; coded alive.', 'Yes'),
    ('Hand', 489, None, 'Abolished Oct. 1941 (succeeded by Lend-Lease Administration)', '615', 'Abolished before 1943; coded alive.', 'Yes'),
    ('Hand', 490, None, 'Renamed Board of War Communications, June 15, 1942', '64', 'Renamed before 1943; coded alive.', 'Yes'),
    ('Hand', 500, None, 'Functions to National Housing Agency, Feb. 24, 1942', '606', 'Abolished before 1943; coded alive.', 'Yes'),
    ('Hand', 300, None, 'Terminated Oct. 13, 1942', '609', 'Abolished before 1943; coded alive.', 'Yes'),
    ('Hand', 356, None, 'Title expired June 22, 1942', '615', 'Abolished before 1943; coded alive.', 'Yes'),
    ('Hand', 265, None, 'Replaced in 1940 by the Bureau of the Public Debt (Fiscal Service)', '206', 'Replaced in 1940; coded alive.', 'Yes'),
    ('Hand', 17, None, 'Commerce lists the Civil Aeronautics Board and Administrator of Civil Aeronautics, not the Authority', '368', 'Authority split up in 1940; coded alive.', 'Yes'),
    # --- Hand-entered: correct for June 1943, wrong for December 1943 ---
    ('Hand', 493, None, 'Consolidated into Foreign Economic Administration, Sept. 25, 1943', '68', 'Ended after June 1943.', 'No'),
    ('Hand', 512, None, 'Terminated July 15, 1943 (became Office of Economic Warfare, then FEA)', '598, 68', 'Ended after June 1943.', 'No'),
    ('Hand', 374, None, 'Abolished effective Aug. 31, 1943', '57', 'Ended after June 1943.', 'No'),
    ('Hand', 363, None, 'Not in the Interior list; Bituminous Coal Act expired Aug. 1943', '306; 599-600', 'Ended after June 1943.', 'No'),
    ('Hand', 375, None, 'Liquidation ordered by act of July 12, 1943', '618; 420', 'Ended after June 1943.', 'No'),
    ('Hand', 317, None, 'Functions transferred to the Federal Works Administrator, June 30, 1943', '445; 611', 'Ended around June 1943.', 'Possibly'),
    ('Hand', 130, None, 'Not in the list of State divisions and offices', '182', 'May have existed in June 1943; check that edition.', 'Possibly'),
    ('Hand', 151, None, 'Not in the list of State divisions and offices', '182', 'May have existed in June 1943; check that edition.', 'Possibly'),
    ('Hand', 208, None, 'Not in the Interior list of bureaus and offices', '306', 'May have existed in June 1943; check that edition.', 'Possibly'),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = '1943 reading errors'
head = ['Side', 'Row', 'Agency', 'My 1943 coding', 'Hand 1943 coding', 'What the Dec. 1943 manual shows',
        'Page(s)', 'Explanation', 'Wrong in June 1943 edition too?']
ws.append(head)
for side, r, mine_override, shows, pages, why, june in E:
    name = str(mine_ws.cell(r, 1).value).strip()
    hr = hand_rows.get(re.sub(r'[^a-z]', '', name.lower()))
    mine = mine_override if mine_override else code(mine_ws, r)
    hand = code(hand_ws, hr) if hr else '(not in hand sheet)'
    ws.append([side, r, name, mine, hand, shows, pages, why, june])

bold = Font(bold=True, color='FFFFFF')
fill = PatternFill('solid', fgColor='305496')
for c in ws[1]:
    c.font, c.fill = bold, fill
    c.alignment = Alignment(wrap_text=True, vertical='center')
widths = [8, 6, 42, 34, 34, 52, 12, 44, 14]
for i, w in enumerate(widths):
    ws.column_dimensions[chr(65 + i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top')
    if row[0].value == 'Mine':
        for c in row:
            c.fill = PatternFill('solid', fgColor='FCE4D6')
ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions

s = wb.create_sheet('Summary')
n_m = sum(1 for e in E if e[0] == 'Mine')
n_h = sum(1 for e in E if e[0] == 'Hand')
n_hy = sum(1 for e in E if e[0] == 'Hand' and e[6] == 'Yes')
for line in [
    ['1943 reading errors, both sides'],
    ['Source', 'United States Government Manual, Winter 1943-44 (Dec. 1, 1943); pages are printed page numbers'],
    ['My errors', n_m, 'all fixed'],
    ['Hand-entered errors', n_h, f'{n_hy} wrong in either 1943 edition; the rest depend on whether the June 1943 manual was used'],
    ['Not included', 'Convention differences (parent level, renamed rows, international bodies), rows not filled in, and liquidation judgment calls (CCC, WPA)'],
]:
    s.append(line)
s['A1'].font = Font(bold=True, size=13)
s.column_dimensions['A'].width = 22
s.column_dimensions['B'].width = 80
wb.save(OUT)
print('saved', OUT, '| mine', n_m, '| hand', n_h, f'({n_hy} wrong in both editions)')
