"""Text tools for a U.S. Government Manual PDF (searching only — decide structure from page images).

Usage (python manual_text.py <command> ...):
  extract   <pdf> <out.txt>                 extract text, one "=====PAGE n" block per PDF page
  contents  <txt>                           print the contents page and the PDF-to-printed page offset
  appendix  <txt> <year>                    Appendix A entries that mention <year> or <year-1> (abolished/transferred)
  heads     <txt> [first] [last]            running section headers by PDF page (maps the book)
  lists     <txt> [first] [last]            "principal bureaus ... as follows:" unit lists (finds pages to render)
  find      <txt> <term> [term ...]         PDF pages + context for each term (regex, case-insensitive)
  presence  <txt> <sheet.xlsx> <col>        live rows in that year column whose names are not found in the text

PDF page numbers are what render_pages.py takes; printed page = PDF page - offset.
"""
import collections, re, sys


def load(txt):
    t = open(txt, encoding='utf-8').read()
    parts = re.split(r'=====PAGE (\d+)\n', t)[1:]
    return {int(parts[i]): parts[i + 1] for i in range(0, len(parts), 2)}


def flat(s):
    return re.sub(r'\s+', ' ', re.sub(r'-\n', '', s))


def norm(s):
    s = s.lower().replace('’', "'")
    s = re.sub(r"[^a-z0-9' ]", ' ', s)
    return re.sub(r'\s+', ' ', s).strip()


def offset(P):
    c = collections.Counter()
    for k, v in P.items():
        L = [l.strip() for l in v.split('\n') if l.strip()]
        for l in L[:2] + L[-2:]:
            m = re.match(r'^(\d{1,3})(\s|$)', l) or re.search(r'(\s|^)(\d{1,3})$', l)
            if m:
                n = int(m.group(1) if m.re.pattern.startswith('^') else m.group(2))
                if 5 < n < 750:
                    c[k - n] += 1
    return c.most_common(3)


def cmd_extract(pdf, out):
    from pypdf import PdfReader
    r = PdfReader(pdf)
    with open(out, 'w', encoding='utf-8') as f:
        for i, p in enumerate(r.pages):
            try:
                t = p.extract_text() or ''
            except Exception as e:  # damaged page: keep going
                t = f'[ERR {e}]'
            f.write(f'=====PAGE {i + 1}\n{t}\n')
    print(out, len(r.pages), 'pages')


def cmd_contents(txt):
    P = load(txt)
    for k in sorted(P):
        if k < 30 and 'CONTENTS' in P[k].upper()[:300]:
            print('PDF page', k, ':', flat(P[k])[:1500])
            break
    print('PDF minus printed page (most common first):', offset(P))
    if 'charts' in ' '.join(P.get(k, '') for k in range(1, 30)).lower() and 'discontinued' in ' '.join(P.get(k, '') for k in range(1, 30)).lower():
        print('NOTE: organization charts are discontinued in this edition — use unit lists and officials pages.')


def cmd_appendix(txt, year):
    P = load(txt)
    y = int(year)
    starts = [k for k in sorted(P) if k > 400 and P[k].strip().upper().startswith('APPENDIX A')]
    if not starts:
        starts = [k for k in sorted(P) if k > 400 and 'ABOLISHED' in P[k][:400].upper()]
    a = starts[0]
    b = next((k for k in sorted(P) if k > a and 'APPENDIX B' in P[k][:200].upper()), a + 40)
    text = ' '.join(flat(P[k]) for k in range(a, b))
    ents = re.split(r"(?<=[.)]) (?=[A-Z][A-Z,'() &.-]{3,90}\.\s?[-—–])", text)
    pat = re.compile(rf'{y}|{y - 1}')
    print(f'Appendix A on PDF pages {a}-{b - 1}; entries mentioning {y - 1} or {y}:\n')
    for e in ents:
        if pat.search(e):
            sents = [s for s in re.split(r'(?<=\.) ', e) if pat.search(s)]
            print('*', e[:100], '::', ' '.join(sents)[:600], '\n')
    print('Two-column pages extract badly — render the appendix pages as images when an entry matters.')


def cmd_heads(txt, first='1', last='9999'):
    P = load(txt)
    prev = None
    for k in sorted(P):
        if not int(first) <= k <= int(last):
            continue
        L = [l.strip() for l in P[k].split('\n') if l.strip()]
        if not L:
            continue
        h = re.sub(r'\s+\d+$', '', re.sub(r'^\d+\s*', '', L[0]))
        let = [c for c in h if c.isalpha()]
        if 'MANUAL' in h.upper() or len(let) < 6 or sum(c.isupper() for c in let) / len(let) < 0.85:
            continue
        if h != prev:
            print(k, h)
            prev = h


def cmd_lists(txt, first='40', last='700'):
    P = load(txt)
    pat = re.compile(r'(follows ?:|are as follows|consists of the following|comprises the following|following offices[^.:]{0,40}:)', re.I)
    for k in sorted(P):
        if not int(first) <= k <= int(last):
            continue
        F = flat(P[k])
        for m in pat.finditer(F):
            seg = F[m.end():m.end() + 500]
            if len(seg) > 40:
                print(f'[PDF {k}] ...{F[max(0, m.start() - 90):m.start()]}|| {seg}\n')


def cmd_find(txt, *terms):
    P = load(txt)
    F = {k: flat(v) for k, v in P.items()}
    for t in terms:
        hits = [(k, m.start()) for k in sorted(F) for m in re.finditer(t, F[k], re.I)]
        print(f'## {t}: PDF pages {sorted(set(k for k, _ in hits))[:20]}')
        for k, s in hits[:3]:
            print('   ', k, F[k][max(0, s - 150):s + 150])


def cmd_presence(txt, sheet, col):
    import openpyxl
    P = load(txt)
    T = norm(' '.join(flat(P[k]) for k in P))
    ws = openpyxl.load_workbook(sheet).worksheets[0]
    c = int(col)
    miss = []
    for r in range(3, ws.max_row + 1):
        n = ws.cell(r, 1).value
        v = [str(ws.cell(r, x).value).strip() for x in range(c, c + 5)]
        if not n or not (v[0] == 'Yes' or v[1] == 'Yes' or v[4] == 'Yes'):
            continue
        k = norm(re.sub(r'^the ', '', str(n).strip(), flags=re.I).rstrip(' *'))
        k = re.sub(r"'s?$", '', re.sub(r'\s*\([^)]*\)$', '', k)).strip()
        if k and k not in T:
            miss.append(f'{r}:{n}')
    print(len(miss), 'live rows not found (often misspelled row names or renamed units — check each):')
    print(' | '.join(miss))


if __name__ == '__main__':
    cmds = {'extract': cmd_extract, 'contents': cmd_contents, 'appendix': cmd_appendix, 'heads': cmd_heads,
            'lists': cmd_lists, 'find': cmd_find, 'presence': cmd_presence}
    if len(sys.argv) < 2 or sys.argv[1] not in cmds:
        print(__doc__)
        sys.exit(1)
    cmds[sys.argv[1]](*sys.argv[2:])
