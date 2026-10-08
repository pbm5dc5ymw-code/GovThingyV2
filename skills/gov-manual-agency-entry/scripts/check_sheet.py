"""Consistency checks for one year of the agency spreadsheet.

Usage: python check_sheet.py <sheet.xlsx> <year> [<previous_sheet.xlsx>]
Checks:
  - counts by category for the year
  - every parent named in the year is itself a live row that year (exact name match)
  - duplicate agency names whose codings disagree
  - if a previous sheet is given: no cells in earlier-year columns changed
  - number of hidden rows
"""
import collections, sys
import openpyxl

LIVE = ('Cabinet', 'Child of cabinet', 'Child of other', 'Other')


def cat(ws, r, c):
    v = [str(ws.cell(r, x).value).strip() if ws.cell(r, x).value is not None else '' for x in range(c, c + 5)]
    if not any(v): return 'blank'
    if v[0] == 'Dead': return 'Dead'
    if v[0] == 'Yes': return 'Cabinet'
    if v[1] == 'Yes': return 'Child of cabinet' if v[2] not in ('None', '') else 'Child of other'
    if v[4] == 'Yes': return 'Other'
    return 'Not listed yet'


path, year = sys.argv[1], int(sys.argv[2])
ws = openpyxl.load_workbook(path).worksheets[0]
hdr = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value}
C = hdr[f'{year}_cabinet_agency']
rows = [r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value]
print(year, dict(collections.Counter(cat(ws, r, C) for r in rows)))

names = {}
for r in rows:
    names.setdefault(str(ws.cell(r, 1).value).strip(), []).append(r)
bad = []
for r in rows:
    for x in (ws.cell(r, C + 2).value, ws.cell(r, C + 3).value):
        x = str(x).strip() if x else ''
        if x and x not in ('None', 'Dead', 'No'):
            if not any(cat(ws, pr, C) in LIVE for pr in names.get(x, [])):
                bad.append(f'row {r}: parent "{x}" is not a live row in {year}')
print('Parent problems:', bad or 'none')

dups = [(n, rs) for n, rs in names.items() if len(rs) > 1 and len({tuple(ws.cell(r, C + i).value for i in range(5)) for r in rs}) > 1]
print('Same name in several rows with different codings (often different departments — check, not necessarily an error):', dups or 'none')

if len(sys.argv) > 3:
    prev = openpyxl.load_workbook(sys.argv[3]).worksheets[0]
    changed = [(r, c) for r in range(1, prev.max_row + 1) for c in range(3, C) if prev.cell(r, c).value != ws.cell(r, c).value]
    print('Earlier-year cells changed:', len(changed), changed[:10])

print('Hidden rows:', sum(1 for r in rows if ws.row_dimensions[r].hidden))
